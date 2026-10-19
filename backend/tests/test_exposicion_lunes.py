"""Seis pruebas de integración reales para presentar sin Docker."""

import io
import re

import pytest
from jose import jwt
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from app.core.config import settings


@pytest.fixture
def client(isolated_client):
    return isolated_client


@pytest.fixture
def auth_headers(client):
    response = client.post(
        "/api/v1/auth/token",
        data={"username": "admin", "password": "Admin@12345"},
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_01_health_and_status(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "healthy"


def test_02_auth_login_jwt(client):
    response = client.post(
        "/api/v1/auth/token",
        data={"username": "admin", "password": "Admin@12345"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["token_type"] == "bearer"
    claims = jwt.decode(body["access_token"], settings.SECRET_KEY, algorithms=[settings.ALGORITHM])
    assert claims["sub"] == "admin"
    assert claims["type"] == "access"
    assert claims["exp"] > claims["iat"]
    rejected = client.post(
        "/api/v1/auth/token",
        data={"username": "admin", "password": "clave-equivocada"},
    )
    assert rejected.status_code == 401


def test_03_ai_drone_detection_real(client, auth_headers):
    image = Image.new("RGB", (320, 240), (135, 206, 235))
    ImageDraw.Draw(image).rectangle((120, 40, 190, 90), fill=(10, 10, 10))
    buffer = io.BytesIO()
    image.save(buffer, format="JPEG")
    response = client.post(
        "/api/v1/ai/analyze",
        headers=auth_headers,
        files={"file": ("drone.jpg", buffer.getvalue(), "image/jpeg")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "success"
    assert isinstance(body["engine"], str) and body["engine"]
    detection = body["detection"]
    assert detection["is_drone"] is True
    assert 0 <= detection["confidence"] <= 1
    box = detection["bounding_box"]
    assert 0 <= box["x_min"] < 320
    assert 0 <= box["y_min"] < 240
    assert 0 < box["width"] <= 320
    assert 0 < box["height"] <= 240


def test_04_ai_invalid_file_rejected(client, auth_headers):
    response = client.post(
        "/api/v1/ai/analyze",
        headers=auth_headers,
        files={"file": ("falso.txt", b"Esto no es una imagen", "text/plain")},
    )
    assert response.status_code == 400
    assert response.json()["detail"]


def test_05_report_pdf_sha256(client, auth_headers):
    response = client.get("/api/v1/reports/export/pdf", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF-")
    assert b"%%EOF" in response.content[-1024:]
    assert re.fullmatch(r"[0-9a-f]{64}", response.headers["X-Report-SHA256"])
    assert ".pdf" in response.headers["content-disposition"]


def test_06_report_excel_export(client, auth_headers):
    response = client.get("/api/v1/reports/export/excel", headers=auth_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert ".xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(response.content))
    try:
        assert "Telemetría e Incursiones" in workbook.sheetnames
        assert "Resumen y Auditoría" in workbook.sheetnames
        assert workbook["Telemetría e Incursiones"].max_row >= 5
    finally:
        workbook.close()
