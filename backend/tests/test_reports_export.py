"""
AeroGuard — Tests para Exportación Física de Reportes (RF-024 a RF-027 / CODE-02)
Valida los criterios de aceptación Gherkin de HU-02:
- Escenario 3: Descarga de reporte táctico en PDF binario real con hash SHA-256.
- Escenario 4: Exportación de telemetría histórica a Excel (.xlsx) válido.
- Control de acceso RBAC: Analyst y Admin tienen acceso a reportes, Operador denegado.
"""

import io
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.core.security import create_access_token


@pytest.fixture
def client():
    return TestClient(app)


def test_export_pdf_report_success(client: TestClient):
    """Escenario 3: Analista solicita exportar reporte PDF y recibe documento binario válido."""
    token = create_access_token({"sub": "analyst_user", "role": "analyst"})

    response = client.get(
        "/api/v1/reports/export/pdf",
        headers={"Authorization": f"Bearer {token}"},
        params={"type": "detections"},
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.pdf"')

    # Verificar que el PDF empieza con la firma binaria de PDF (%PDF)
    content = response.content
    assert len(content) > 1000
    assert content.startswith(b"%PDF")

    # Verificar que tiene el header de hash criptográfico SHA-256
    sha256 = response.headers.get("X-Report-SHA256")
    assert sha256 is not None
    assert len(sha256) == 64


def test_export_excel_report_success(client: TestClient):
    """Escenario 4: Analista solicita exportar reporte Excel y recibe libro .xlsx válido."""
    token = create_access_token({"sub": "analyst_user", "role": "analyst"})

    response = client.get(
        "/api/v1/reports/export/excel",
        headers={"Authorization": f"Bearer {token}"},
        params={"type": "detections"},
    )

    assert response.status_code == 200
    assert "openxmlformats" in response.headers["content-type"]
    assert "attachment; filename=" in response.headers["content-disposition"]
    assert response.headers["content-disposition"].endswith('.xlsx"')

    content = response.content
    assert len(content) > 2000

    # Cargar el binario con openpyxl para verificar que es un Excel válido y bien formado
    wb = load_workbook(io.BytesIO(content))
    assert "Telemetría e Incursiones" in wb.sheetnames
    assert "Resumen y Auditoría" in wb.sheetnames

    ws1 = wb["Telemetría e Incursiones"]
    assert ws1["A1"].value.startswith("AEROGUARD C-UAS")
    # Verificar que tiene filas de incursiones
    assert ws1.max_row >= 5


def test_reports_rbac_operator_denied(client: TestClient):
    """Prueba RBAC: Rol operator no tiene permiso reports:read."""
    token = create_access_token({"sub": "operator_user", "role": "operator"})

    response = client.get(
        "/api/v1/reports/export/pdf",
        headers={"Authorization": f"Bearer {token}"},
    )

    assert response.status_code == 403
    assert "No tienes permiso para: reports:read" in response.json().get("detail", "")


def test_reports_post_export_unified(client: TestClient):
    """Prueba endpoint POST unificado /api/v1/reports/export."""
    token = create_access_token({"sub": "admin_user", "role": "admin"})

    # Exportar PDF vía POST
    res_pdf = client.post(
        "/api/v1/reports/export",
        headers={"Authorization": f"Bearer {token}"},
        json={"format": "pdf", "type": "detections"},
    )
    assert res_pdf.status_code == 200
    assert res_pdf.content.startswith(b"%PDF")

    # Exportar Excel vía POST
    res_xl = client.post(
        "/api/v1/reports/export",
        headers={"Authorization": f"Bearer {token}"},
        json={"format": "excel", "type": "detections"},
    )
    assert res_xl.status_code == 200
    assert "openxmlformats" in res_xl.headers["content-type"]
